"""Optional end-to-end check against the running app and the real installed model.
Uses invented input, verifies downloads, and removes its output documents.
"""
from io import BytesIO
import os
import time
from urllib.parse import quote
import httpx
import pymupdf
from docx import Document

BASE = os.environ.get('BASE_URL', 'http://127.0.0.1:3001') + '/api/service'
SENSITIVITY = int(os.environ.get('SENSITIVITY', '50'))
TEXT = 'Please contact Alice Smith at alice.smith@example.com. Alice Smith is our account contact.'


def verify(client, identifier):
    for _ in range(600):
        row = client.get(BASE + '/guest/current').json()['document']
        assert row and row['id'] == identifier, 'Current result changed unexpectedly'
        if row['status'] == 'failed':
            raise AssertionError(row['error'])
        if row['status'] == 'complete':
            break
        time.sleep(1)
    else:
        raise TimeoutError('Model job did not finish within ten minutes.')
    assert row['sensitivity'] == SENSITIVITY, row
    assert row['filename'] == 'Smoke résumé report.' + row['source_type'], row
    assert row['counts'].get('private_person') == 2, row
    assert row['counts'].get('private_email') == 1, row
    for ext in ('txt', 'docx', 'pdf'):
        response = client.get(f'{BASE}/guest/documents/{identifier}/download/{ext}')
        response.raise_for_status()
        if ext == 'txt':
            text = response.text
        elif ext == 'docx':
            text = '\n'.join(p.text for p in Document(BytesIO(response.content)).paragraphs)
        else:
            with pymupdf.open(stream=response.content, filetype='pdf') as pdf:
                text = '\n'.join(p.get_text() for p in pdf)
        if row['mode'] == 'redact':
            assert '******' in text, (ext, text)
        assert 'Alice Smith' not in text and 'alice.smith@example.com' not in text, (ext, text)
    review = client.get(f'{BASE}/guest/documents/{identifier}/review')
    review.raise_for_status()
    assert 'Alice Smith' not in review.text and 'alice.smith@example.com' not in review.text
    report = review.json()['scan_report']
    assert report['total_detections'] == 3 and report['sensitivity'] == SENSITIVITY
    assert report['ocr_performed'] is False
    revealed = client.get(f'{BASE}/guest/documents/{identifier}/revealed-detections')
    revealed.raise_for_status()
    assert revealed.headers['cache-control'] == 'no-store'
    assert set(revealed.json()) == {'values'}
    assert len(revealed.json()['values']) == report['total_detections']
    for detection in revealed.json()['values']:
        assert set(detection) == {'category', 'occurrence', 'original'}
        expected = 'Alice Smith' if detection['category'] == 'private_person' else 'alice.smith@example.com'
        assert detection['original'] == expected
    with httpx.Client(timeout=30) as stranger:
        stranger.get(BASE + '/guest/current').raise_for_status()
        assert stranger.get(f'{BASE}/guest/documents/{identifier}/download/txt').status_code == 404
        assert stranger.get(f'{BASE}/guest/documents/{identifier}/revealed-detections').status_code == 404
    return row['counts']


with httpx.Client(timeout=30, limits=httpx.Limits(max_keepalive_connections=0)) as client:
    current = client.get(BASE + '/guest/current')
    current.raise_for_status()
    client.headers['X-CSRF-Token'] = current.json()['csrf_token']
    for suffix, mode in (('.docx', 'redact'), ('.pdf', 'redact'), ('.docx', 'placeholder'), ('.pdf', 'synthetic')):
        if suffix == '.docx':
            doc = Document(); doc.add_paragraph(TEXT)
            buffer = BytesIO(); doc.save(buffer); payload = buffer.getvalue()
        else:
            pdf = pymupdf.open(); page = pdf.new_page()
            page.insert_textbox(pymupdf.Rect(50, 50, 540, 200), TEXT)
            payload = pdf.tobytes(); pdf.close()
        response = client.post(BASE + f'/guest/documents?type={suffix}&mode={mode}&sensitivity={SENSITIVITY}', content=payload,
                               headers={'X-Document-Name': quote('Smoke résumé report' + suffix)})
        response.raise_for_status()
        identifier = response.json()['id']
        try:
            counts = verify(client, identifier)
            print(f'{suffix} / {mode}: real inference and all three exports passed; {counts}', flush=True)
        finally:
            # Reset also invalidates a still-running job if an assertion/network check fails.
            deleted = client.delete(BASE + '/guest/current')
            deleted.raise_for_status()
            assert client.get(f'{BASE}/guest/documents/{identifier}/download/txt').status_code == 404
            current = client.get(BASE + '/guest/current')
            current.raise_for_status()
            client.headers['X-CSRF-Token'] = current.json()['csrf_token']
    print('Real-model smoke passed. Test documents deleted.')
