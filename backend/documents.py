"""Extract text and create fresh files; never copy original containers or metadata."""
from collections import Counter
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile
import html
import secrets
import pymupdf
from docx import Document
from docx.table import Table

LABELS = {
    'private_person': 'Name', 'private_address': 'Address',
    'private_email': 'Email', 'private_phone': 'Phone Number',
    'private_date': 'Date', 'private_url': 'URL',
    'account_number': 'Account Number', 'secret': 'Secret',
}


def blocks(container):
    for block in container.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                yield ' | '.join('\n'.join(blocks(cell)) for cell in row.cells)
        else:
            yield block.text


def extract(data: bytes, suffix: str) -> str:
    if suffix == '.pdf':
        with pymupdf.open(stream=data, filetype='pdf') as pdf:
            if pdf.is_encrypted:
                raise ValueError('Password-protected PDFs are not supported.')
            if len(pdf) > 250:
                raise ValueError('Please use a document with at most 250 pages.')
            pages = []
            for page in pdf:
                text = page.get_text(sort=True)
                if not text.strip() and page.get_images():
                    raise ValueError('This PDF contains a scanned page. Run OCR locally first.')
                pages.append(text)
            text = '\n\n'.join(pages)
    elif suffix == '.docx':
        try:
            with ZipFile(BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError('The expanded Word document exceeds 100 MB.')
                if 'word/document.xml' not in archive.namelist():
                    raise ValueError('This is not a valid DOCX document.')
        except BadZipFile:
            raise ValueError('This is not a valid DOCX document.') from None
        doc = Document(BytesIO(data))
        parts = list(blocks(doc))
        for section in doc.sections:
            for part in (section.header, section.first_page_header, section.even_page_header,
                         section.footer, section.first_page_footer, section.even_page_footer):
                parts.extend(blocks(part))
        text = '\n'.join(parts)
    else:
        raise ValueError('Upload a PDF or DOCX file. Convert legacy .doc files to .docx first.')
    if not text.strip():
        raise ValueError('No readable text was found. Scanned documents need OCR first.')
    if len(text) > 200_000:
        raise ValueError('Extracted text exceeds 200,000 characters. Split the document first.')
    return text


def transform(result, mode: str):
    # OPF offsets refer to result.text, which may differ from the original tokenizer input.
    text = result.text
    cursor, pieces, counts, mapping = 0, [], Counter(), {}
    nonce = secrets.token_hex(3)
    for span in sorted(result.detected_spans, key=lambda s: (s.start, s.end)):
        if span.label not in LABELS or not 0 <= cursor <= span.start < span.end <= len(text):
            raise ValueError('The model returned unsupported or overlapping spans; no output was saved.')
        original = text[span.start:span.end]
        key = (span.label, original.casefold())
        if key not in mapping:
            n = len(mapping) + 1
            dummy = {
                'private_person': f'Alex Example {n}',
                'private_address': f'{n} Example Street, Sampletown',
                'private_email': f'person{n}.{nonce}@example.invalid',
                'private_phone': f'+1 202 555 {100 + (n % 100):04d}',
                'private_date': f'2000-01-{1 + n % 28:02d}',
                'private_url': f'https://example.invalid/reference/{nonce}/{n}',
                'account_number': f'DEMO-{nonce.upper()}-{n:06d}',
                'secret': f'SYNTHETIC-{nonce}-{n}',
            }
            mapping[key] = f'[{LABELS[span.label]}]' if mode == 'placeholder' else dummy[span.label]
        pieces.extend((text[cursor:span.start], mapping[key]))
        cursor = span.end
        counts[span.label] += 1
    pieces.append(text[cursor:])
    return ''.join(pieces), dict(counts)


def export_files(text: str, folder: Path):
    folder.mkdir(parents=True, exist_ok=False)
    (folder / 'sanitized.txt').write_text(text, encoding='utf-8')
    doc = Document()
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    for line in text.split('\n'):
        doc.add_paragraph(line)
    doc.save(folder / 'sanitized.docx')
    story = pymupdf.Story(html='<html><body>' + ''.join(
        '<p>' + html.escape(line) + '</p>' for line in text.split('\n')
    ) + '</body></html>', user_css='body { font-family: sans-serif; font-size: 11pt; } p { overflow-wrap: anywhere; }')
    with pymupdf.DocumentWriter(str(folder / 'sanitized.pdf')) as writer:
        more = True
        pages = 0
        while more:
            pages += 1
            if pages > 1000:
                raise ValueError('The rebuilt PDF exceeded the output page limit.')
            device = writer.begin_page(pymupdf.Rect(0, 0, 595, 842))
            more, _ = story.place(pymupdf.Rect(45, 45, 550, 797))
            story.draw(device)
            writer.end_page()
