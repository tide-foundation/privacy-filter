from io import BytesIO
from types import SimpleNamespace
import time
from fastapi.testclient import TestClient
from docx import Document
from backend import app as service


class Detector:
    _runtime = True
    def redact(self, text):
        start = text.index('Alice')
        return SimpleNamespace(text=text, detected_spans=[SimpleNamespace(start=start, end=start+5, label='private_person')], warning=None)


def test_upload_download_delete_and_origin_protection(tmp_path, monkeypatch):
    monkeypatch.setattr(service, 'ROOT', tmp_path)
    monkeypatch.setattr(service, 'model', Detector())
    original_find = service.importlib.util.find_spec
    monkeypatch.setattr(service.importlib.util, 'find_spec', lambda name: True if name == 'opf' else original_find(name))
    doc = Document(); doc.add_paragraph('Hello Alice.'); data = BytesIO(); doc.save(data)
    with TestClient(service.app, base_url='http://127.0.0.1:8000') as client:
        client.headers['X-CSRF-Token'] = client.get('/api/service/guest/current').json()['csrf_token']
        assert client.post('/api/service/guest/documents?type=.docx', content=data.getvalue(), headers={'origin': 'https://evil.example'}).status_code == 403
        assert client.post('/api/service/guest/documents?type=.doc', content=b'invalid').status_code == 400
        response = client.post('/api/service/guest/documents?type=.docx&mode=placeholder', content=data.getvalue())
        assert response.status_code == 202
        identifier = response.json()['id']
        for _ in range(100):
            items = [client.get('/api/service/guest/current').json()['document']]
            if items[0]['status'] in {'complete', 'failed'}: break
            time.sleep(.02)
        assert items[0]['status'] == 'complete', items
        assert items[0]['counts'] == {'private_person': 1}
        assert items[0]['layout_preserved'] == 1
        assert 'Alice' not in str(items)
        assert client.get(f'/api/service/guest/documents/{identifier}/preview').json()['text'].startswith('Hello [Name].')
        for ext in ('pdf', 'docx', 'txt'):
            assert client.get(f'/api/service/guest/documents/{identifier}/download/{ext}').status_code == 200
        assert client.get('/api/service/guest/documents/not-a-uuid/preview').status_code == 404
        assert client.delete(f'/api/service/guest/documents/{identifier}').status_code == 200
        assert not (tmp_path / 'guest' / identifier).exists()
        assert client.get('/api/service/guest/current').json()['document'] is None
        assert client.get(f'/api/service/guest/documents/{identifier}/download/pdf').status_code == 404


def test_changed_tokenizer_text_fails_without_saving_outputs(tmp_path, monkeypatch):
    class MismatchDetector:
        _runtime = True
        def redact(self, text):
            return SimpleNamespace(text='Changed ' + text, detected_spans=[], warning='Mismatch')
    monkeypatch.setattr(service, 'ROOT', tmp_path)
    monkeypatch.setattr(service, 'model', MismatchDetector())
    original_find = service.importlib.util.find_spec
    monkeypatch.setattr(service.importlib.util, 'find_spec', lambda name: True if name == 'opf' else original_find(name))
    doc = Document(); doc.add_paragraph('Hello Alice.'); data = BytesIO(); doc.save(data)
    with TestClient(service.app, base_url='http://127.0.0.1:8000') as client:
        client.headers['X-CSRF-Token'] = client.get('/api/service/guest/current').json()['csrf_token']
        response = client.post('/api/service/guest/documents?type=.docx', content=data.getvalue())
        identifier = response.json()['id']
        for _ in range(100):
            items = [client.get('/api/service/guest/current').json()['document']]
            if items[0]['status'] == 'failed': break
            time.sleep(.02)
        assert items[0]['status'] == 'failed'
        assert 'tokenization' in items[0]['error']
        assert items[0]['source_type'] == 'docx'
        assert not (tmp_path / 'guest' / identifier).exists()
        assert client.get(f'/api/service/guest/documents/{identifier}/download/docx').status_code == 404


def test_pdf_fit_failure_completes_with_clean_downloads(tmp_path, monkeypatch):
    import pymupdf
    class ShortNameDetector:
        _runtime = True
        def redact(self, text):
            start = text.index('Li')
            return SimpleNamespace(text=text, detected_spans=[SimpleNamespace(start=start, end=start+2, label='private_person')], warning=None)
    monkeypatch.setattr(service, 'ROOT', tmp_path)
    monkeypatch.setattr(service, 'model', ShortNameDetector())
    original_find = service.importlib.util.find_spec
    monkeypatch.setattr(service.importlib.util, 'find_spec', lambda name: True if name == 'opf' else original_find(name))
    pdf = pymupdf.open(); page = pdf.new_page(); page.insert_text((50, 100), 'Li', fontsize=12)
    data = pdf.tobytes(); pdf.close()
    with TestClient(service.app, base_url='http://127.0.0.1:8000') as client:
        client.headers['X-CSRF-Token'] = client.get('/api/service/guest/current').json()['csrf_token']
        for mode in ('redact', 'placeholder', 'synthetic'):
            response = client.post(f'/api/service/guest/documents?type=.pdf&mode={mode}', content=data)
            assert response.status_code == 202
            identifier = response.json()['id']
            for _ in range(100):
                row = client.get('/api/service/guest/current').json()['document']
                if row['status'] in {'complete', 'failed'}: break
                time.sleep(.02)
            assert row['status'] == 'complete', row
            assert row['layout_preserved'] == 0
            assert row['source_type'] == 'pdf'
            assert row['error'] is None
            assert row['warning'] == 'Original layout unavailable; clean rewrite used.'
            assert row['counts'] == {'private_person': 1}
            replacement = {'redact': '******', 'placeholder': '[Name]', 'synthetic': 'Alex Example 1'}[mode]
            for ext in ('txt', 'pdf', 'docx'):
                response = client.get(f'/api/service/guest/documents/{identifier}/download/{ext}')
                assert response.status_code == 200
                if ext == 'txt': text = response.text
                elif ext == 'docx': text = '\n'.join(p.text for p in Document(BytesIO(response.content)).paragraphs)
                else:
                    with pymupdf.open(stream=response.content, filetype='pdf') as output:
                        text = ''.join(p.get_text() for p in output)
                assert replacement in text and 'Li' not in text
            assert not list((tmp_path / 'guest' / identifier).glob('.native*'))


def test_sensitivity_validation_and_job_setting(tmp_path, monkeypatch):
    monkeypatch.setattr(service, 'ROOT', tmp_path)
    monkeypatch.setattr(service, 'model', Detector())
    original_find = service.importlib.util.find_spec
    monkeypatch.setattr(service.importlib.util, 'find_spec', lambda name: True if name == 'opf' else original_find(name))
    seen = []
    def calibrated(model, text, sensitivity, cache_dir):
        seen.append(sensitivity)
        return model.redact(text)
    monkeypatch.setattr(service, 'redact', calibrated)
    doc = Document(); doc.add_paragraph('Hello Alice.'); data = BytesIO(); doc.save(data)
    with TestClient(service.app, base_url='http://127.0.0.1:8000') as client:
        client.headers['X-CSRF-Token'] = client.get('/api/service/guest/current').json()['csrf_token']
        for invalid in ('-5', '105', '51', '75.5', 'nan', ''):
            assert client.post(f'/api/service/guest/documents?type=.docx&sensitivity={invalid}', content=data.getvalue()).status_code == 400
        assert client.get('/api/service/guest/current').json()['document'] is None
        for level in (75, 25):
            response = client.post(f'/api/service/guest/documents?type=.docx&sensitivity={level}', content=data.getvalue())
            assert response.status_code == 202
            identifier = response.json()['id']
            for _ in range(100):
                row = client.get('/api/service/guest/current').json()['document']
                if row['status'] in {'complete', 'failed'}: break
                time.sleep(.02)
            assert row['status'] == 'complete', row
            assert row['sensitivity'] == level
        assert seen == [75, 25]
