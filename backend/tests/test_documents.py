from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile
import pytest
import pymupdf
from docx import Document
from backend.documents import extract, transform, export_files


def result(text, entries):
    return SimpleNamespace(text=text, detected_spans=[SimpleNamespace(start=s, end=e, label=label) for s, e, label in entries], warning=None)


def test_unicode_offsets_and_repeated_synthetic_values():
    text = '🙂 Alice met Alice on Tuesday.'
    spans = [(2, 7, 'private_person'), (12, 17, 'private_person'), (21, 28, 'private_date')]
    output, counts = transform(result(text, spans), 'placeholder')
    assert output == '🙂 [Name] met [Name] on [Date].'
    output, _ = transform(result(text, spans), 'synthetic')
    assert output.count('Alex Example 1') == 2
    assert 'Alice' not in output and 'Tuesday' not in output
    assert counts == {'private_person': 2, 'private_date': 1}


def test_invalid_spans_fail_closed():
    with pytest.raises(ValueError):
        transform(result('Alice', [(0, 4, 'private_person'), (3, 5, 'secret')]), 'placeholder')


def test_docx_table_header_extraction_and_clean_exports(tmp_path):
    doc = Document()
    doc.add_paragraph('Hello Alice')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'alice@example.com'
    doc.sections[0].header.paragraphs[0].text = 'Private header'
    doc.core_properties.author = 'Original sensitive author'
    source = BytesIO(); doc.save(source)
    text = extract(source.getvalue(), '.docx')
    assert 'Hello Alice' in text and 'alice@example.com' in text and 'Private header' in text
    target = tmp_path / 'outputs'
    export_files('Hello [Name]\n[Email]', target)
    assert (target / 'sanitized.txt').read_text() == 'Hello [Name]\n[Email]'
    with ZipFile(target / 'sanitized.docx') as archive:
        assert b'Original sensitive author' not in archive.read('docProps/core.xml')
        assert b'Alice' not in archive.read('word/document.xml')
    with pymupdf.open(target / 'sanitized.pdf') as pdf:
        output = ''.join(p.get_text() for p in pdf)
        assert '[Name]' in output and '[Email]' in output and 'Alice' not in output


def test_reject_invalid_or_empty_files():
    with pytest.raises(ValueError): extract(b'not a docx', '.docx')
    with pytest.raises(ValueError): extract(b'old word', '.doc')
    doc = Document(); source = BytesIO(); doc.save(source)
    with pytest.raises(ValueError, match='No readable text'): extract(source.getvalue(), '.docx')


def test_scanned_pdf_is_rejected():
    doc = pymupdf.open(); page = doc.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, (0, 0, 10, 10), False)
    pix.clear_with(255)
    page.insert_image(pymupdf.Rect(0, 0, 100, 100), pixmap=pix)
    with pytest.raises(ValueError, match='scanned page'): extract(doc.tobytes(), '.pdf')
    doc.close()


@pytest.mark.parametrize('suffix', ['.pdf', '.docx'])
@pytest.mark.parametrize('error_type', [ValueError, RuntimeError])
def test_partial_native_export_keeps_clean_rewrite(tmp_path, suffix, error_type):
    from backend.documents import export_with_fallback
    def failed_save(path, edits):
        path.write_bytes(b'partial invalid native output')
        raise error_type('Native export failed after writing')
    source = SimpleNamespace(save=failed_save)
    folder = tmp_path / 'outputs'
    assert not export_with_fallback('Hello [Name].', folder, source, suffix, [])
    assert sorted(p.name for p in folder.iterdir()) == ['sanitized.docx', 'sanitized.pdf', 'sanitized.txt']
    assert (folder / 'sanitized.txt').read_text() == 'Hello [Name].'
    assert Document(folder / 'sanitized.docx').paragraphs[0].text == 'Hello [Name].'
    with pymupdf.open(folder / 'sanitized.pdf') as pdf:
        assert 'Hello [Name].' in pdf[0].get_text()


@pytest.mark.parametrize('label', ['private_person', 'private_address', 'private_email', 'private_phone', 'private_date', 'private_url', 'account_number', 'secret'])
def test_redact_masks_do_not_reveal_length_or_format(label):
    values = ['Li', 'A very long sensitive value', '2026-01-02', '2 January 2026', '01/02/26']
    expected = '**/**/**' if label == 'private_date' else '******'
    for value in values:
        text = f'Before {value} after.'
        output, counts = transform(result(text, [(7, 7 + len(value), label)]), 'redact')
        assert output == f'Before {expected} after.'
        assert counts == {label: 1}
